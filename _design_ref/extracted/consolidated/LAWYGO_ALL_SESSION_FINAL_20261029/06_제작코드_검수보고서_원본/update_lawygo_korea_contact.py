from pathlib import Path
from docx import Document
from docx.shared import Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re, json, html
r=Path('/mnt/data/LAWYGO_PARTNERSHIP_SITE')
address='서울특별시 서초구 서초중앙로22길 47, INS-154'
addr_en='INS-154, 47 Seochojungang-ro 22-gil, Seocho-gu, Seoul, Republic of Korea'
phone='0505-248-2468'; fax='0505-284-2468'; email='shinkang888@gmail.com'; main_email='sk@terrabridge.us'
contact='''<section class="korea-office" aria-label="LAWYGO 국내 파트너십 연락처"><div class="wrap korea-office-layout"><div><span class="eyebrow">LAWYGO / KOREA LIAISON</span><h2>국내 법무법인 협력 창구</h2><p>국내 영업·제휴 문의는 <strong>대한에이전트협회</strong>를 통해 안내합니다. 미국 법인과의 업무협약은 담당자 검토 및 계약 당사자 확인 후 개별 협의합니다.</p></div><div class="office-info"><p><b>영업·제휴 창구</b> 대한에이전트협회 (Korea Agent Association)</p><p><b>주소</b> 서울특별시 서초구 서초중앙로22길 47, INS-154</p><p><b>전화</b> <a href="tel:05052482468">0505-248-2468</a> <span class="sep">·</span> <b>팩스</b> 0505-284-2468</p><p><b>국내 문의</b> <a href="mailto:shinkang888@gmail.com">shinkang888@gmail.com</a></p><p><b>국제 제휴</b> <a href="mailto:sk@terrabridge.us">sk@terrabridge.us</a></p><p class="note">대한에이전트협회는 국내 연락 창구로 표기하며, TERRABRIDGE CAPITAL LLC의 법적 지점 또는 계약 당사자라는 의미가 아닙니다.</p></div></div></section>'''
css='''\n/* Korean LAWYGO partnership business contact */
.korea-office{background:#f2f0e9;border-top:1px solid #e0d9ce;border-bottom:1px solid #dedbd3;padding:48px 0;color:#19304a}
.korea-office-layout{display:grid;grid-template-columns:1fr 1.3fr;gap:7vw;align-items:start}
.korea-office h2{font-family:Georgia,'Noto Serif KR',serif;font-size:clamp(25px,3vw,38px);line-height:1.22;margin:10px 0 16px}
.korea-office p{font-size:13px;color:#46586b;line-height:1.9;margin:6px 0}
.korea-office .office-info{background:#fff;border:1px solid #e5dfd4;padding:22px 26px;border-radius:10px}
.korea-office .office-info b{color:#142b40;margin-right:10px}
.korea-office .office-info a{color:#155274;text-decoration:underline;text-underline-offset:3px}
.korea-office .office-info .note{border-top:1px solid #eee;padding-top:10px;font-size:10px;color:#7b858e}
.korea-office .sep{opacity:.5;margin:0 7px}
.korea-office .eyebrow{font-size:11px;color:#a48452;letter-spacing:.18em}
.footer-korea-contact{margin:16px 0;padding:15px 0;border-top:1px solid #67748955;font-size:12px;color:#b6c1c9;line-height:1.85}
.footer-korea-contact strong{color:#e2d0aa}
.footer-korea-contact a{color:#d5c9b5;text-decoration:underline;text-underline-offset:3px}
.footer-korea-contact .footer-domain{font-size:11px;opacity:.8}
@media(max-width:700px){.korea-office-layout{grid-template-columns:1fr;gap:20px}.korea-office{padding:30px 0}.korea-office .office-info{padding:18px}.footer-korea-contact{font-size:11px}}
'''
(r/'assets'/'korea-contact.css').write_text(css,'utf8')
footer_contact='''<div class="footer-korea-contact"><strong>대한민국 영업·제휴 창구 | 대한에이전트협회</strong><br>서울특별시 서초구 서초중앙로22길 47, INS-154<br>전화 <a href="tel:05052482468">0505-248-2468</a> · 팩스 0505-284-2468 · <a href="mailto:shinkang888@gmail.com">shinkang888@gmail.com</a><br>미국 파트너십 <a href="mailto:sk@terrabridge.us">sk@terrabridge.us</a> <span class="footer-domain"> | 공식 도메인 lawygo.com</span></div>'''
count=0
for f in r.rglob('*.html'):
    if f.name in ['partnership-inquiry.html','mou.html']:continue
    s=f.read_text('utf8')
    rel='../' if f.parent!=r else ''
    if 'korea-contact.css' not in s:
        s=s.replace('</head>',f'<link rel="stylesheet" href="{rel}assets/korea-contact.css"></head>',1)
    if '<footer class="footer">' in s and 'footer-korea-contact' not in s:
        # add above footer bottom
        s=s.replace('<div class="footer-bottom">',footer_contact+'<div class="footer-bottom">',1)
    if f.name=='partnership.html' and f.parent==r and 'class="korea-office"' not in s:
        s=s.replace('</main>',contact+'</main>',1)
    if f.name=='contact.html' and f.parent==r and 'class="korea-office"' not in s:
        s=s.replace('</main>',contact+'</main>',1)
    if f.name=='about.html' and f.parent==r and 'class="korea-office"' not in s:
        s=s.replace('</main>',contact+'</main>',1)
    # Canonical for public landing, keep offline local navigation paths working
    if f.name=='partnership.html' and 'rel="canonical"' not in s:
        s=s.replace('</head>','<link rel="canonical" href="https://lawygo.com/partnership"></head>',1)
    if f.name=='index.html' and f.parent==r and 'rel="canonical"' not in s:
        s=s.replace('</head>','<link rel="canonical" href="https://lawygo.com/"></head>',1)
    if f.name=='contact.html' and f.parent==r and 'rel="canonical"' not in s:
        s=s.replace('</head>','<link rel="canonical" href="https://lawygo.com/contact"></head>',1)
    f.write_text(s,'utf8');count+=1

# add Korea service contact in MOU online sample, making KAA non-contractual
f=r/'mou.html';s=f.read_text('utf8')
s=s.replace('갑과 을은 아래와 같이 상호 협력에 관한 기본 의향을 확인한다.','갑과 을은 아래와 같이 상호 협력에 관한 기본 의향을 확인한다.<br><br><b>국내 제휴 연락 창구 (계약 당사자 아님)</b>: 대한에이전트협회<br>'+address+'<br>전화 '+phone+' · 팩스 '+fax+'<br>국내 이메일 '+email+' · 국제 문의 '+main_email+'<br>공식 사이트: lawygo.com')
s=s.replace('<p class="small">본 문서는 협의용','<p class="small">한국 제휴 연락처: 대한에이전트협회 · '+address+' · 전화 '+phone+' · 팩스 '+fax+' · '+email+'<br>본 문서는 협의용')
f.write_text(s,'utf8')

# popup contact
f=r/'partnership-inquiry.html';s=f.read_text('utf8')
s=s.replace('수신: <a href="mailto:sk@terrabridge.us">sk@terrabridge.us</a>','국제 제휴 수신: <a href="mailto:sk@terrabridge.us">sk@terrabridge.us</a> · 국내 창구: <a href="mailto:shinkang888@gmail.com">shinkang888@gmail.com</a>')
s=s.replace('</form></main>','</form><p class="small"><b>국내 영업·제휴 창구 — 대한에이전트협회</b><br>서울특별시 서초구 서초중앙로22길 47, INS-154<br>전화 0505-248-2468 · 팩스 0505-284-2468 · 공식 도메인 lawygo.com</p></main>')
f.write_text(s,'utf8')

# Site domain plan documented, do not point menus at unconfigured DNS.
domains={
 'root':{'domain':'lawygo.com','purpose':'LAWYGO 브랜드·법률경제 플랫폼','site_relative':'index.html'},
 'journal':{'domain':'journal.lawygo.com','purpose':'법률경제 저널·월 2회 무료 뉴스레터','site_relative':'journal/index.html'},
 'wiki':{'domain':'wiki.lawygo.com','purpose':'법률백과·판례정보','site_relative':'wiki/index.html'},
 'note':{'domain':'note.lawygo.com','purpose':'AI 문서 분석·서식작성','site_relative':'note/index.html'},
 'pro':{'domain':'pro.lawygo.com','purpose':'법률 송무관리·무료 베타','site_relative':'pro/index.html'},
 'pan':{'domain':'pan.lawygo.com','purpose':'변호사 큐레이션·마케팅','site_relative':'pan/index.html'},
 'partner':{'domain':'lawygo.com/partnership','purpose':'법무법인 MOU 제휴 신청','site_relative':'partnership.html'},
 'payment':{'domain':'lawygo.com/payment','purpose':'구독상품·결제 안내','site_relative':'payment.html'} }
(r/'DOMAIN_MAP.json').write_text(json.dumps(domains,ensure_ascii=False,indent=2),'utf8')
(r/'DOMAIN_DEPLOYMENT_PLAN.md').write_text('''# LAWYGO 브랜드 도메인 및 배포설계 (2026-10-29)

**공식 기준 도메인:** `lawygo.com` (기존 사이트 기반). 아래 세부 서비스 주소는 **배포 제안**이며 실제 DNS/SSL/라우팅 연결을 확인한 상태는 아닙니다.

|서비스|권장 URL|동봉된 HTML 시안 위치|
|---|---|---|
'''+''.join(f'|{k.upper()}|`https://{v["domain"]}`|`{v["site_relative"]}`|\n' for k,v in domains.items())+'''

## 배포 체크리스트
1. 호스팅(Vercel 또는 기존 lawygo.com 서버)에 `lawygo.com`, `journal.lawygo.com`, `wiki.lawygo.com`, `note.lawygo.com`, `pro.lawygo.com`, `pan.lawygo.com` 도메인을 하나씩 등록하고 DNS CNAME/A 검증 후 SSL을 활성화합니다.
2. **각 도메인이 연결되기 전까지 임의의 subdomain 외부 링크로 즉시 치환하지 않습니다.** ZIP은 파일 상대경로로 모든 페이지 간 이동이 가능하도록 유지했습니다.
3. 기존 lawygo 프로젝트로 실제 송무를 제공하는 `pro.lawygo.com`은 기존 Google OAuth callback 허용목록, 세션 쿠키 Domain·SameSite·Secure 설정, 프로젝트별 DB/멀티테넌트 권한을 검수한 후 연결합니다.
4. LAWYGO 결제는 기존 `shinkangteam/lexid`의 PR #14 (`/lawygo/payment`)로 설계됐지만 **운영 병합·키 검증·판매 승인 전**입니다. 새로운 `lawygo.com/payment`를 공개할 때에는 결제 서버와 로그인을 보안 경계 내에서 연결하고 CORS/CSRF/Stripe Webhook URL·성공/취소/계정 접근을 교차 도메인 기준 재검증합니다.
5. Partnership: `lawygo.com/partnership` (정적 `partnership.html`을 rewrite), MOU는 `lawygo.com/mou`, 간편 문의창은 `lawygo.com/partnership-inquiry` rewrite를 권장합니다. 현재 폼은 메일 앱 `mailto:`를 열 뿐 자동 접수/DB 저장 기능이 아닙니다.
6. 개인정보처리방침 및 서비스 약관을 각 결제·회원가입·데이터 수집처리 범위에 맞게 확정하고 제휴 폼의 수신·보관 기간을 고지합니다.

## 대외 명칭·국내 창구
- 한국 협력·영업 창구: **대한에이전트협회 (Korea Agent Association)**
- 주소: 서울특별시 서초구 서초중앙로22길 47, INS-154 (명함 영문: `INS-154, 47, Seochojungang-ro 22-gil, Seocho-gu, Seoul, Republic of Korea`)
- 전화: 0505-248-2468 / 팩스: 0505-284-2468
- 국내 문의: shinkang888@gmail.com
- 미국 제휴 이메일: sk@terrabridge.us
- **법적 주의:** 국내 창구 표기는 TERRABRIDGE CAPITAL LLC의 등록된 한국 지점·계약 대리권을 증명하지 않습니다. MOU 서명법인과 권한은 서명 시 검증합니다.
''','utf8')

# Update brand doc and partnership guide
with (r/'PARTNERSHIP_GUIDE.md').open('a',encoding='utf8') as fh:
    fh.write('\n## 한국내 연락처 및 도메인 (2026-10-29 수정)\n- 국내 창구: 대한에이전트협회 · '+address+' · '+phone+' · 팩스 '+fax+' · '+email+'\n- 국제 제휴 문의: '+main_email+'\n- 서비스 도메인 계획: `DOMAIN_DEPLOYMENT_PLAN.md` / `DOMAIN_MAP.json` (실제 DNS 연결 전)\n')
with (r/'BRAND_ARCHITECTURE.md').open('a',encoding='utf8') as fh:
    fh.write('\n## LAWYGO 정식 도메인 구조 제안\n'+''.join(f'- {k}: https://{v["domain"]}\n' for k,v in domains.items()))
# Update Word MOU with contact box after party-identification table before first clause
p=r/'documents'/'LAWYGO_법무법인_MOU_약정서.docx';doc=Document(p)
par=doc.add_paragraph('국내 연락 창구 (본 MOU 계약 당사자 아님): 대한에이전트협회')
par.style=doc.styles['Normal'];par.paragraph_format.space_before=Pt(7);par.paragraph_format.space_after=Pt(0)
for x in par.runs:x.bold=True;x.font.color.rgb=RGBColor(27,59,82);x.font.size=Pt(8.5)
par2=doc.add_paragraph(f'{address}  |  전화 {phone}  |  팩스 {fax}\n국내 이메일 {email}  |  국제 제휴 {main_email}  |  공식 도메인 lawygo.com')
par2.paragraph_format.space_after=Pt(5)
for x in par2.runs:x.font.size=Pt(8)
intro=None
for z in doc.paragraphs:
 if z.text.startswith('갑과 을은 각자의 독립적인'):
  intro=z
  break
if intro is not None:
 intro._p.addprevious(par._p)
 intro._p.addprevious(par2._p)
# Provide contact line at end as reference for signature form
end=doc.add_paragraph();end.paragraph_format.space_before=Pt(7)
x=end.add_run('연락처 (한국 제휴 창구): 대한에이전트협회 | '+address+' | '+phone+' | 팩스 '+fax+' | '+email)
x.font.size=Pt(7.2);x.font.color.rgb=RGBColor(105,113,119)
doc.save(p)
print('Patched',count,'HTML files, site and MOU',p.stat().st_size)
