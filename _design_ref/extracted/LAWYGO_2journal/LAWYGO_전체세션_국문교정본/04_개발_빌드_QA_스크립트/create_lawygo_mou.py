from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path('/mnt/data/LAWYGO_PARTNERSHIP_SITE'); out=root/'documents'/'LAWYGO_법무법인_MOU_약정서.docx'; clauses=json.loads((root/'MOU_CLAUSES.json').read_text('utf8'))
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.3);sec.header_distance=Cm(.8);sec.footer_distance=Cm(.85)
styles=doc.styles;normal=styles['Normal']; normal.font.name='Noto Sans CJK KR';normal.font.size=Pt(9.7);normal.font.color.rgb=RGBColor(35,49,62);normal.paragraph_format.line_spacing=1.42;normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',20),('Heading 1',12.1)]:
 s=styles[name];s.font.name='Noto Sans CJK KR';s.font.size=Pt(size);s.font.color.rgb=RGBColor(20,43,64);s.font.bold=True
 s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(6)
header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('LAWYGO  /  PARTNERSHIP MOU  ·  협의용 초안');r.bold=True;r.font.size=Pt(7.5);r.font.color.rgb=RGBColor(138,119,91)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('LAWYGO · 한국 파트너 법무법인 업무협약 초안  |  법률·직역 검토 및 당사자 확인 후 서명');r.font.size=Pt(7);r.font.color.rgb=RGBColor(125,135,145)
par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
par.add_run().add_picture(str(root/'assets'/'lawygo-logo-1024.png'),width=Cm(2.4))
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(7)
r=p.add_run('MEMORANDUM OF UNDERSTANDING');r.font.size=Pt(9);r.bold=True;r.font.color.rgb=RGBColor(170,136,76)
p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('업무협약서 (MOU)')
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(12)
r=p.add_run('LAWYGO × 대한민국 파트너 법무법인');r.font.color.rgb=RGBColor(112,121,128);r.font.size=Pt(10)
# parties table
t=doc.add_table(rows=2,cols=2);t.autofit=False;t.columns[0].width=Cm(2.4);t.columns[1].width=Cm(13.8)
for i,(a,b) in enumerate([('갑','TERRABRIDGE CAPITAL LLC  (LAWYGO 관련 사업운영 주체 · 정식 서명 전 확인)'),('을','____________________________  (법무법인 / 법률사무소)')]):
 t.cell(i,0).text=a;t.cell(i,1).text=b
 for cell in t.rows[i].cells:
  cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  cell._tc.get_or_add_tcPr().append(OxmlElement('w:shd'))
  cell._tc.tcPr[-1].set(qn('w:fill'),'F4F1EA' if i==0 else 'FAFAF8')
  for p in cell.paragraphs:
   for run in p.runs:run.font.size=Pt(9);run.bold=i==0
p=doc.add_paragraph('갑과 을은 각자의 독립적인 권한과 전문성을 존중하며 다음과 같이 협력의 기본 방향을 정한다.');p.paragraph_format.space_before=Pt(10)
for title,body in clauses:
 h=doc.add_paragraph(style='Heading 1');h.add_run(title);h.paragraph_format.keep_with_next=True
 for piece in body.split(' ② '):
  p=doc.add_paragraph();p.paragraph_format.widow_control=True
  p.paragraph_format.keep_together=True
  if piece!=body.split(' ② ')[0]:piece='② '+piece
  p.add_run(piece)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(19);p.add_run('체결일  ______년  ______월  ______일')
t=doc.add_table(rows=4,cols=2);t.autofit=False;t.columns[0].width=Cm(8.1);t.columns[1].width=Cm(8.1)
for col,name in enumerate(['갑 · TERRABRIDGE CAPITAL LLC','을 · 파트너 법무법인']):
 t.cell(0,col).text=name
 t.cell(1,col).text='대표권자: ______________________' if col==0 else '대표변호사: ______________________'
 t.cell(2,col).text='주소: __________________________' if col==0 else '법무법인명: ______________________'
 t.cell(3,col).text='서명: __________________________' if col==0 else '서명: __________________________'
 for row in range(4):
  cell=t.cell(row,col)
  for p in cell.paragraphs:
   p.paragraph_format.space_after=Pt(7)
   for run in p.runs:
    run.font.size=Pt(9)
    if row==0:run.bold=True;run.font.color.rgb=RGBColor(20,47,67)
p=doc.add_paragraph();p.paragraph_format.space_before=Pt(10)
r=p.add_run('※ 사용 안내: 본 문서는 협상용 초안입니다. 법인 등록·서명권자, 광고 규정 준수, 데이터 이전, 책임 범위 및 비용 조건은 쌍방 검토 후 확정하십시오.');r.italic=True;r.font.size=Pt(8);r.font.color.rgb=RGBColor(130,107,75)
doc.save(out)
print(out)
